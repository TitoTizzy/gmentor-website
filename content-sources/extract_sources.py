from __future__ import annotations

import json
import shutil
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent
PYTHON = Path(r"C:\Users\ouhha\.cache\codex-runtimes\codex-primary-runtime\dependencies\python\python.exe")
PDFTOPPM = Path(r"C:\Users\ouhha\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe")

SOURCES = [
    {
        "id": "usa-resume",
        "market": "usa",
        "kind": "docx",
        "path": Path(r"C:\Users\ouhha\OneDrive\Documents\Marie Gaelle Mentor\Document Word\Marie Gaelle Mentor Resume.docx"),
    },
    {
        "id": "usa-portfolio",
        "market": "usa",
        "kind": "docx",
        "path": Path(r"C:\Users\ouhha\Downloads\portfolio Mme Mentor (1).docx"),
    },
    {
        "id": "haiti-livret",
        "market": "haiti",
        "kind": "pdf",
        "path": Path(r"C:\Users\ouhha\OneDrive\Documents\Marie Gaelle Mentor\Les Entreprises\MGM\Livret MGM\Livret MGM-Haiti.pdf"),
    },
    {
        "id": "mixed-portfolio",
        "market": "mixed",
        "kind": "pdf",
        "path": Path(r"C:\Users\ouhha\OneDrive\Documents\Marie Gaelle Mentor\Les Entreprises\MGM\Stuff\portfolio Mme Mentor_Haiti et USA.pdf"),
    },
]


def ensure_clean_dir(path: Path) -> None:
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True)


def image_info(path: Path) -> dict:
    try:
        with Image.open(path) as image:
            return {
                "file": path.name,
                "width": image.width,
                "height": image.height,
                "format": image.format,
            }
    except Exception as exc:
        return {"file": path.name, "error": str(exc)}


def extract_docx(source: dict, out_dir: Path) -> dict:
    media_dir = out_dir / "media"
    media_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source["path"]) as archive:
        for name in archive.namelist():
            if name.startswith("word/media/") and not name.endswith("/"):
                target = media_dir / Path(name).name
                with archive.open(name) as source_file, target.open("wb") as output_file:
                    shutil.copyfileobj(source_file, output_file)

    document = Document(source["path"])
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            style = paragraph.style.name if paragraph.style else ""
            blocks.append(f"[{style}] {text}" if style else text)
    for table_index, table in enumerate(document.tables, start=1):
        blocks.append(f"\n[TABLE {table_index}]")
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip().replace("\n", " / ") for cell in row.cells))
    (out_dir / "text.txt").write_text("\n".join(blocks), encoding="utf-8")

    images = [image_info(path) for path in sorted(media_dir.iterdir()) if path.is_file()]
    (out_dir / "media-index.json").write_text(json.dumps(images, indent=2), encoding="utf-8")
    return {"paragraphs": len(blocks), "media": len(images), "images": images}


def make_contact_sheets(page_dir: Path, output_dir: Path, columns: int = 4) -> list[str]:
    page_files = sorted(page_dir.glob("page-*.png"), key=lambda path: int(path.stem.split("-")[-1]))
    output_dir.mkdir(parents=True, exist_ok=True)
    sheets: list[str] = []
    for start in range(0, len(page_files), 12):
        batch = page_files[start : start + 12]
        thumbs = []
        for page_file in batch:
            with Image.open(page_file) as image:
                thumb = image.convert("RGB")
                thumb.thumbnail((360, 480))
                thumbs.append((page_file, thumb.copy()))
        rows = (len(thumbs) + columns - 1) // columns
        sheet = Image.new("RGB", (columns * 380, rows * 520), "white")
        draw = ImageDraw.Draw(sheet)
        for index, (page_file, thumb) in enumerate(thumbs):
            x = (index % columns) * 380 + 10
            y = (index // columns) * 520 + 28
            sheet.paste(thumb, (x, y))
            draw.text((x, 7 + (index // columns) * 520), page_file.stem.replace("page-", "Page "), fill="black", font=ImageFont.load_default())
        end = start + len(batch)
        name = f"pages-{start + 1:03d}-{end:03d}.jpg"
        sheet.save(output_dir / name, quality=88, optimize=True)
        sheets.append(name)
    return sheets


def extract_pdf(source: dict, out_dir: Path) -> dict:
    page_dir = out_dir / "pages"
    ensure_clean_dir(page_dir)
    reader = PdfReader(source["path"])
    text_parts: list[str] = []
    page_text: list[dict] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        text_parts.append(f"\n===== PAGE {index} =====\n{text}")
        page_text.append({"page": index, "characters": len(text), "preview": text[:240]})
    (out_dir / "text-by-page.txt").write_text("\n".join(text_parts), encoding="utf-8")
    (out_dir / "page-text-index.json").write_text(json.dumps(page_text, indent=2, ensure_ascii=False), encoding="utf-8")

    prefix = page_dir / "page"
    subprocess.run(
        [str(PDFTOPPM), "-png", "-r", "96", str(source["path"]), str(prefix)],
        check=True,
    )
    rendered = sorted(page_dir.glob("page-*.png"))
    sheets = make_contact_sheets(page_dir, out_dir / "contact-sheets")
    return {"pages": len(reader.pages), "rendered_pages": len(rendered), "contact_sheets": sheets}


def main() -> None:
    manifest = {"separation_rule": "Never move content between markets without explicit source evidence.", "sources": []}
    for source in SOURCES:
        out_dir = ROOT / source["market"] / source["id"]
        ensure_clean_dir(out_dir)
        result = extract_docx(source, out_dir) if source["kind"] == "docx" else extract_pdf(source, out_dir)
        manifest["sources"].append(
            {
                "id": source["id"],
                "market": source["market"],
                "kind": source["kind"],
                "source": str(source["path"]),
                "output": str(out_dir.relative_to(ROOT)),
                **result,
            }
        )
    (ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(manifest, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
